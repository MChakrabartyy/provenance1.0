"""Manager C: Excel with a header row on line 3 and quantities as strings like "1,200"; Side column gives sign."""
from __future__ import annotations

from pathlib import Path
from typing import Iterator

from openpyxl import load_workbook

from provenance.ingestion.adapters.base import dec, ts
from provenance.models import RawTrade, Trade

HEADER_ROW = 3


class ManagerCAdapter:
    manager_id = "C"

    def parse(self, path: Path) -> Iterator[RawTrade]:
        ws = load_workbook(path, read_only=True, data_only=True).active
        rows = ws.iter_rows(min_row=HEADER_ROW, values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows)]
        for i, values in enumerate(rows, start=HEADER_ROW + 1):
            if all(v is None for v in values):
                continue
            yield RawTrade(self.manager_id, path.name, dict(zip(header, values)), i)

    def to_standard(self, raw: RawTrade) -> Trade:
        r = raw.row
        qty = dec(r.get("Quantity"))
        side = str(r.get("Side") or "BUY").strip().upper()
        if qty is not None and side in ("SELL", "SHORT", "S"):
            qty = -abs(qty)
        return Trade(
            manager_id=self.manager_id,
            symbol=str(r.get("Security") or "").strip().upper(),
            qty=qty,
            price=dec(r.get("Price")),
            trade_time=ts(r.get("Executed"), ("%Y-%m-%d %H:%M",)),
            external_trade_id=str(r.get("Ref") or f"{raw.source_file}:{raw.row_number}"),
            source_file=raw.source_file,
            raw={k: (str(v) if v is not None else None) for k, v in r.items()},
        )
